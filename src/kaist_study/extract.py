from __future__ import annotations

from pathlib import Path
from typing import Iterator


def extract(path: str, *, ocr: bool = False) -> list[tuple[str, str]]:
    """Return location/content pairs; extraction failure is separate from download failure."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return [(f"page {i + 1}", page.extract_text() or "") for i, page in enumerate(PdfReader(path).pages)]
    if suffix == ".pptx":
        from pptx import Presentation
        out: list[tuple[str, str]] = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text]
            notes = getattr(slide, "notes_slide", None)
            if notes:
                texts.extend(shape.text for shape in notes.notes_text_frame.paragraphs if shape.text)
            out.append((f"slide {i}", "\n".join(texts)))
        return out
    raise ValueError(f"unsupported extraction format: {suffix or 'unknown'}; legacy .ppt requires conversion")

